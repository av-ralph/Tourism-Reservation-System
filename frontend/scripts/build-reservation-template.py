"""Refresh the fillable reservation template from its existing official form layout."""

from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


public_dir = Path(__file__).resolve().parents[1] / "public"
template_path = public_dir / "documents/reservation-template.docx"
source = template_path.read_bytes()
document = Document(BytesIO(source))

def add_masthead(document):
    header = document.sections[0].header
    if header.tables:
        return
    masthead = header.add_table(rows=1, cols=3, width=Inches(7.35))
    masthead.autofit = False
    for cell, width in zip(masthead.rows[0].cells, (Inches(1.1), Inches(5.15), Inches(1.1))):
        cell.width = width
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cell.paragraphs[0].paragraph_format.space_after = Pt(0)
    masthead.cell(0, 0).paragraphs[0].add_run().add_picture(str(public_dir / "cec.png"), width=Inches(0.82))
    center = masthead.cell(0, 1).paragraphs[0]
    for index, (line, size) in enumerate((("Cristal e-College", 12), ("Academic Division", 10), ("Hospitality Management", 10))):
        if index:
            center.add_run().add_break()
        run = center.add_run(line)
        run.bold = True
        run.font.name = "Times New Roman"
        run.font.size = Pt(size)
    masthead.cell(0, 2).paragraphs[0].add_run().add_picture(str(Path(__file__).resolve().parent / "assets/hm-form-logo.png"), width=Inches(0.82))
    # The empty header paragraph came from the official form; use its space for the masthead.
    header._element.remove(header.paragraphs[0]._element)


add_masthead(document)

area = document.paragraphs[11]
area.clear()
area.style = "Body Text"
area.alignment = WD_ALIGN_PARAGRAPH.CENTER
area.paragraph_format.space_before = Pt(7)
area.paragraph_format.space_after = Pt(3)

areas = (
    ("BAR", "Bar Area"),
    ("DINING", "Dining Area"),
    ("FRONT_OFFICE", "Front Office Area"),
    ("HOT_KITCHEN", "Hot Kitchen"),
    ("COLD_KITCHEN", "Cold Kitchen"),
    ("LAUNDRY", "Laundry Area"),
    ("MINI_FUNCTION", "Mini Function Room"),
    ("HARMONIA", "Harmonia Mini-Hotel Room"),
    ("LUNARA", "Lunara Mini-Hotel Room"),
    ("COMPUTER", "Computer Laboratory MH 302"),
)
for index, (key, label) in enumerate(areas):
    if index in (3, 6, 8):
        area.add_run().add_break()
    elif index:
        area.add_run("    ")
    box = area.add_run("{{CHECK_" + key + "}}")
    box.font.name = "Segoe UI Symbol"
    box.font.size = Pt(11)
    name = area.add_run(" " + label)
    name.font.name = "Times New Roman"
    name.font.size = Pt(10)

selected = document.paragraphs[12]
selected.clear()
selected.alignment = WD_ALIGN_PARAGRAPH.CENTER
selected.paragraph_format.space_after = Pt(6)
label = selected.add_run("Selected space: ")
label.bold = True
value = selected.add_run("{{FACILITY}}")
value.bold = True

for table in document.tables[:2]:
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.bold = True

approval = document.tables[-1]
for cell in approval.rows[0].cells:
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP


def add_line(cell, text, *, bold=False, size=9, first=False, alignment=WD_ALIGN_PARAGRAPH.LEFT):
    paragraph = cell.paragraphs[0] if first else cell.add_paragraph()
    paragraph.alignment = alignment
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold
    return paragraph


def signature_line(cell, value):
    paragraph = add_line(cell, value)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    paragraph.paragraph_format.line_spacing = Pt(48)
    paragraph.paragraph_format.space_after = Pt(4)
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "5")
    bottom.set(qn("w:color"), "666666")
    border.append(bottom)
    paragraph._p.get_or_add_pPr().append(border)


reserved, noted = approval.rows[0].cells
for cell, side in ((reserved, "right"), (noted, "left")):
    properties = cell._tc.get_or_add_tcPr()
    for old in properties.findall(qn("w:tcBorders")) + properties.findall(qn("w:tcMar")):
        properties.remove(old)
    margins = OxmlElement("w:tcMar")
    margin = OxmlElement("w:" + side)
    margin.set(qn("w:w"), "180")
    margin.set(qn("w:type"), "dxa")
    margins.append(margin)
    properties.append(margins)

add_line(reserved, "Reserved by", bold=True, size=9, first=True)
signature_line(reserved, "{{SIGNATURE}}")
add_line(reserved, "{{INSTRUCTOR}}", bold=True, size=10)
add_line(reserved, "Signature over printed name", size=8)
add_line(reserved, "Date signed: {{SIGNED_DATE}}", size=8)

add_line(noted, "Noted by", bold=True, size=9, first=True)
signature_line(noted, "{{ADMIN_SIGNATURE}}")
add_line(noted, "{{ADMIN_NAME}}", bold=True, size=10)
add_line(noted, "Laboratory In-Charge · Signature over printed name", size=8)
add_line(noted, "Date signed: {{ADMIN_SIGNED_DATE}}", size=8)

document.save(template_path)
