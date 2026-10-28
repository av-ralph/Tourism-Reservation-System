"""Build the fillable loss/damage acknowledgement template from the supplied official form."""

from io import BytesIO
from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


scripts_dir = Path(__file__).resolve().parent
public_dir = scripts_dir.parent / "public"
source_path = scripts_dir / "assets/loss-damage-source.docx"
template_path = public_dir / "documents/concern-template.docx"
document = Document(source_path)


def logo(path):
    # Print size is under an inch, so keep the downloaded form small.
    image = Image.open(path)
    image.thumbnail((400, 400))
    buffer = BytesIO()
    image.save(buffer, format="PNG", optimize=True)
    buffer.seek(0)
    return buffer


def add_masthead(document):
    header = document.sections[0].header
    header.is_linked_to_previous = False
    masthead = header.add_table(rows=1, cols=3, width=Inches(6.5))
    masthead.autofit = False
    for cell, width in zip(masthead.rows[0].cells, (Inches(1.0), Inches(4.5), Inches(1.0))):
        cell.width = width
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cell.paragraphs[0].paragraph_format.space_after = Pt(0)
    masthead.cell(0, 0).paragraphs[0].add_run().add_picture(logo(public_dir / "cec.png"), width=Inches(0.8))
    center = masthead.cell(0, 1).paragraphs[0]
    for index, (line, size) in enumerate((("Cristal e-College, Inc.", 12), ("Academic Division", 10), ("Hospitality Management", 10))):
        if index:
            center.add_run().add_break()
        run = center.add_run(line)
        run.bold = True
        run.font.name = "Times New Roman"
        run.font.size = Pt(size)
    masthead.cell(0, 2).paragraphs[0].add_run().add_picture(logo(scripts_dir / "assets/hm-form-logo.png"), width=Inches(0.8))
    header._element.remove(header.paragraphs[0]._element)


def set_text(cell, text, *, bold=True, size=10, alignment=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = alignment
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


add_masthead(document)

# The blank lines above the title held space for a masthead that is now in the page header.
for paragraph in document.paragraphs[:4]:
    paragraph._element.getparent().remove(paragraph._element)
paragraphs = document.paragraphs
title = paragraphs[0]
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in title.runs:
    run.text = run.text.strip()

# The official form draws its area boxes as floating shapes; text boxes can be checked when filling.
checkbox_shapes = paragraphs[4]
checkbox_shapes._element.getparent().remove(checkbox_shapes._element)
area = paragraphs[5]
area.clear()
area.alignment = WD_ALIGN_PARAGRAPH.CENTER
area.paragraph_format.space_before = Pt(4)
area.paragraph_format.space_after = Pt(2)
areas = (
    ("BAR", "Bar Area"),
    ("DINING", "Dining Area"),
    ("FRONT_OFFICE", "Front Office Area"),
    ("KITCHEN", "Kitchen Area"),
    ("LAUNDRY", "Laundry Area"),
    ("ROOMS", "Rooms"),
)
for index, (key, label) in enumerate(areas):
    if index:
        area.add_run("   ")
    box = area.add_run("{{CHECK_" + key + "}}")
    box.font.name = "Segoe UI Symbol"
    box.font.size = Pt(11)
    name = area.add_run(" " + label)
    name.font.size = Pt(10)

selected = paragraphs[6]
selected.clear()
selected.alignment = WD_ALIGN_PARAGRAPH.CENTER
selected.paragraph_format.space_before = Pt(4)
selected.paragraph_format.space_after = Pt(6)
label = selected.add_run("Selected space: ")
label.bold = True
label.font.size = Pt(10)
value = selected.add_run("{{FACILITY}}")
value.bold = True
value.font.size = Pt(10)
for paragraph in paragraphs[7:9]:
    paragraph._element.getparent().remove(paragraph._element)

filed, usage, items, signatures = document.tables
# The source floats these tables beside the text; keep them in reading order.
for table in document.tables:
    for position in table._tbl.tblPr.findall(qn("w:tblpPr")):
        table._tbl.tblPr.remove(position)
for cell in items.rows[2].cells[2:5]:
    condition = cell.paragraphs[0]
    condition.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for indent in ("left_indent", "right_indent", "first_line_indent"):
        setattr(condition.paragraph_format, indent, Pt(0))
    for run in condition.runs:
        run.font.size = Pt(8)
        run.text = run.text.strip()
set_text(filed.rows[0].cells[1], "{{FILED_DATE}}")
for row, (left, right) in zip(usage.rows, (("{{USAGE_DATE}}", "{{SUBJECT}}"), ("{{USAGE_TIME}}", "{{PURPOSE}}"))):
    set_text(row.cells[1], left)
    set_text(row.cells[3], right)

# Keep one item row; the browser repeats it for every reported item.
for row in items.rows[4:]:
    row._tr.getparent().remove(row._tr)
for index, key in enumerate(("ITEM_QUANTITY", "ITEM_DESCRIPTION", "ITEM_BROKE", "ITEM_LOST", "ITEM_DAMAGED", "ITEM_REPLACED")):
    set_text(items.rows[3].cells[index], "{{" + key + "}}", size=9, alignment=WD_ALIGN_PARAGRAPH.LEFT if index == 1 else WD_ALIGN_PARAGRAPH.CENTER)

for cell, key in zip(signatures.rows[1].cells, ("STUDENT_NAME", "INSTRUCTOR_NAME", "IN_CHARGE_NAME")):
    set_text(cell, "{{" + key + "}}", alignment=WD_ALIGN_PARAGRAPH.CENTER)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
    cell.paragraphs[0].paragraph_format.space_before = Pt(22)

document.save(template_path)
